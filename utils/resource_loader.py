import csv
from pathlib import Path

from docx import Document as DocxDocument
from langchain_core.documents import Document
from pypdf import PdfReader

RESOURCES_DIR = Path(__file__).resolve().parent.parent / "resources"


def _load_txt(path: Path) -> str:
    return path.read_text(encoding="utf-8", errors="ignore")


def _load_csv(path: Path) -> str:
    with path.open(newline="", encoding="utf-8", errors="ignore") as f:
        rows = list(csv.reader(f))
    return "\n".join(", ".join(row) for row in rows)


def _load_pdf(path: Path) -> str:
    reader = PdfReader(str(path))
    return "\n".join(page.extract_text() or "" for page in reader.pages)


def _load_docx(path: Path) -> str:
    doc = DocxDocument(str(path))
    return "\n".join(paragraph.text for paragraph in doc.paragraphs)


_LOADERS = {
    ".txt": _load_txt,
    ".csv": _load_csv,
    ".pdf": _load_pdf,
    ".docx": _load_docx,
}


def load_resources(directory: str | Path = RESOURCES_DIR) -> list[Document]:
    """Load every supported file in `directory` into a list of Documents."""
    directory = Path(directory)
    documents = []
    for path in sorted(directory.iterdir()):
        if not path.is_file():
            continue
        loader = _LOADERS.get(path.suffix.lower())
        if loader is None:
            continue
        content = loader(path)
        documents.append(Document(page_content=content, metadata={"source": str(path)}))
    return documents
